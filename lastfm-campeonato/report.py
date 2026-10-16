#!/usr/bin/env python3
"""Genera informe.xlsx con el campeonato (matriz mes x artista), clasificación anual,
estadísticas y campeones, a partir de la base SQLite."""
import os, sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font

DB = os.environ.get("DB_PATH", "campeonato.db")

def matriz(con):
    meses = [r[0] for r in con.execute("SELECT DISTINCT mes FROM v_campeonato ORDER BY mes")]
    tot = con.execute("SELECT artista FROM v_estadisticas ORDER BY puntos_totales DESC, victorias DESC").fetchall()
    arts = [r[0] for r in tot]
    pts = {(m, a): p for m, a, p in con.execute("SELECT mes, artista, puntos FROM v_campeonato")}
    return meses, arts, pts

def vista(wb, con, titulo, sql):
    ws = wb.create_sheet(titulo)
    cur = con.execute(sql)
    ws.append([c[0] for c in cur.description])
    for c in ws[1]: c.font = Font(bold=True)
    for row in cur: ws.append(list(row))
    ws.freeze_panes = "A2"

def main():
    con = sqlite3.connect(DB)
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Campeonato")
    meses, arts, pts = matriz(con)
    ws.append([None] + arts)
    for c in ws[1]: c.font = Font(bold=True)
    for m in meses: ws.append([m] + [pts.get((m, a)) for a in arts])
    ws.freeze_panes = "B2"
    vista(wb, con, "Clasificacion anual", "SELECT anio, posicion, artista, puntos, victorias, podios, meses_puntuados FROM v_clasificacion_anual ORDER BY anio, posicion")
    vista(wb, con, "Campeones", "SELECT * FROM v_campeones ORDER BY anio")
    vista(wb, con, "Estadisticas", "SELECT * FROM v_estadisticas ORDER BY victorias DESC, podios DESC, puntos_totales DESC")
    vista(wb, con, "Empates", "SELECT * FROM v_empates")
    wb.save("informe.xlsx"); print("informe.xlsx generado")

if __name__ == "__main__":
    main()
