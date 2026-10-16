#!/usr/bin/env python3
"""Compara los puntos calculados con tu hoja 'Campeonato' del Excel original.
Uso: python validate.py Lastfm.xlsx"""
import re, sqlite3, sys, os, unicodedata
from datetime import datetime
from openpyxl import load_workbook

DB = os.environ.get("DB_PATH", "campeonato.db")

# nombres de tu Excel que difieren de Last.fm (ya normalizados: minúsculas, sin tildes ni signos)
ALIAS = {"lamoda": "lamaravillosaorquestadelalcohol", "triagunlodeamorbizarro": "triangulodeamorbizarro",
         "veitiuno": "veintiuno", "sangujuelas": "sanguijuelasdelguadiana", "sanguijueladelguadiana": "sanguijuelasdelguadiana", "tab": "triangulodeamorbizarro", "siloe": "siloe", "ninosmuntantes": "ninosmutantes"}

def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]", "", s)
    return ALIAS.get(s, s)

def main(path):
    ws = load_workbook(path, data_only=True)["Campeonato"]
    cab = [c.value for c in ws[1]]
    esperado = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], datetime): break
        mes = row[0].strftime("%Y-%m")
        for art, p in zip(cab[1:], row[1:]):
            if art and p: esperado[(mes, norm(art))] = (art, int(p))
    con = sqlite3.connect(DB)
    real = {(m, norm(a)): (a, p) for m, a, p in con.execute("SELECT mes, artista, puntos FROM v_campeonato")}
    meses = {k[0] for k in esperado}
    dif = 0
    for k in sorted(set(esperado) | {k for k in real if k[0] in meses}):
        e, r = esperado.get(k), real.get(k)
        if (e and e[1]) != (r and r[1]):
            dif += 1
            print(f"{k[0]}  {(e or r)[0]:<28} Excel={e[1] if e else '-':>3}  BD={r[1] if r else '-':>3}")
    print(f"\n{len(meses)} meses comparados, {dif} diferencias.")

if __name__ == "__main__":
    main(sys.argv[1])
